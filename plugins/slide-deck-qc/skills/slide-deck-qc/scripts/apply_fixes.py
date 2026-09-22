#!/usr/bin/env python3
"""
Apply the FIX tier to a Syndeio deck. Writes a new file; never edits the
input.

    python apply_fixes.py deck.pptx -o deck-QC.pptx \
        --footer confidential --skip title_drift --log changes.md

--footer  plain | confidential   (required — decides the footer wording)
--skip    a fix kind, repeatable, e.g. --skip whitespace --skip font
--only    restrict to these kinds instead
"""

import argparse
import copy
import re
import sys

from pptx import Presentation
from pptx.util import Inches

EMU = 914400.0
FOOTER_LEFT, FOOTER_TOP = 7.97, 6.74
PAGENUM_LEFT, PAGENUM_TOP = 12.46, 6.74
POS_TOL = 0.05
EDGE_PAD = 0.10

FOOTER_PLAIN = "SYNDEIO BIOSCIENCES"
FOOTER_CONF = "SYNDEIO BIOSCIENCES - CONFIDENTIAL"

ALLOWED_FONTS = {"Aptos", "Aptos Display"}
INHERITED = {None, "", "+mn-lt", "+mj-lt"}
DIVIDER_HINTS = ("transition", "divider", "section")


def inches(v):
    return None if v is None else v / EMU


def walk(shapes, depth=0):
    """Yield (shape, depth). Group children use the group's coordinate
    space, so geometry fixes must only touch depth 0."""
    for sh in shapes:
        yield sh, depth
        if sh.shape_type == 6:
            try:
                for i in walk(sh.shapes, depth + 1):
                    yield i
            except Exception:
                pass


def slide_role(idx, total, layout_name):
    ln = (layout_name or "").lower()
    if idx == 1:
        return "title"
    if idx == total:
        return "back"
    if any(k in ln for k in DIVIDER_HINTS):
        return "divider"
    return "content"


def is_footer(sh, H):
    try:
        return (sh.has_text_frame
                and inches(sh.top) > H - 1.0
                and "SYNDEIO" in sh.text_frame.text.upper())
    except Exception:
        return False


def is_pagenum(sh):
    try:
        return sh.has_text_frame and "slidenum" in sh.text_frame._txBody.xml
    except Exception:
        return False


def find_templates(prs, H):
    """Grab a known-good footer and page-number shape to clone."""
    foot = num = None
    for s in prs.slides:
        for sh in s.shapes:
            if foot is None and is_footer(sh, H):
                foot = sh
            if num is None and is_pagenum(sh):
                num = sh
        if foot is not None and num is not None:
            break
    return foot, num


def clone_onto(slide, shape):
    new = copy.deepcopy(shape._element)
    slide.shapes._spTree.append(new)
    return new


def set_text_keep_format(shape, text):
    """Replace wording without collapsing runs."""
    done = False
    for para in shape.text_frame.paragraphs:
        for i, run in enumerate(para.runs):
            if not done:
                run.text = text
                done = True
            else:
                run.text = ""
    return done


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--footer", choices=["plain", "confidential"],
                    required=True)
    ap.add_argument("--skip", action="append", default=[])
    ap.add_argument("--only", action="append", default=[])
    ap.add_argument("--log")
    a = ap.parse_args()

    want = (lambda k: k in a.only) if a.only else (lambda k: k not in a.skip)
    footer_text = FOOTER_CONF if a.footer == "confidential" else FOOTER_PLAIN

    prs = Presentation(a.deck)
    total = len(prs.slides)
    W, H = inches(prs.slide_width), inches(prs.slide_height)
    changes = []

    def log(slide, kind, msg):
        changes.append((slide, kind, msg))

    foot_tpl, num_tpl = find_templates(prs, H)

    # dominant title position, from content slides
    import collections
    tp = collections.Counter()
    for i, s in enumerate(prs.slides, 1):
        if slide_role(i, total, s.slide_layout.name) != "content":
            continue
        for sh in s.shapes:
            try:
                if sh.is_placeholder and "TITLE" in str(
                        sh.placeholder_format.type):
                    tp[(round(inches(sh.left), 2),
                        round(inches(sh.top), 2))] += 1
            except Exception:
                pass
    dom = tp.most_common(1)[0][0] if tp else None

    for idx, slide in enumerate(prs.slides, 1):
        role = slide_role(idx, total, slide.slide_layout.name)
        wants_chrome = role == "content"
        has_foot = has_num = False

        for sh, depth in list(walk(slide.shapes)):
            try:
                l, t = inches(sh.left), inches(sh.top)
                w, h = inches(sh.width), inches(sh.height)
            except Exception:
                l = t = w = h = None

            # ---- footer ------------------------------------------------
            if is_footer(sh, H):
                has_foot = True
                cur = sh.text_frame.text.strip()
                if want("footer_mixed") and cur != footer_text:
                    if set_text_keep_format(sh, footer_text):
                        log(idx, "footer_mixed",
                            'Footer "%s" -> "%s"' % (cur, footer_text))
                if want("footer_pos") and l is not None and (
                        abs(l - FOOTER_LEFT) > POS_TOL
                        or abs(t - FOOTER_TOP) > POS_TOL):
                    sh.left, sh.top = Inches(FOOTER_LEFT), Inches(FOOTER_TOP)
                    log(idx, "footer_pos",
                        "Footer %.2f/%.2f -> %.2f/%.2f"
                        % (l, t, FOOTER_LEFT, FOOTER_TOP))

            # ---- page number -------------------------------------------
            if is_pagenum(sh):
                has_num = True
                if not wants_chrome and want("pagenum_extra"):
                    sh._element.getparent().remove(sh._element)
                    log(idx, "pagenum_extra",
                        "Removed page number from %s slide" % role)
                    has_num = False
                    continue
                if want("pagenum_pos") and l is not None and (
                        abs(l - PAGENUM_LEFT) > POS_TOL
                        or abs(t - PAGENUM_TOP) > POS_TOL):
                    sh.left, sh.top = Inches(PAGENUM_LEFT), Inches(PAGENUM_TOP)
                    log(idx, "pagenum_pos",
                        "Page number %.2f/%.2f -> %.2f/%.2f"
                        % (l, t, PAGENUM_LEFT, PAGENUM_TOP))

            # ---- fonts --------------------------------------------------
            if want("font") and sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        n = run.font.name
                        if (run.text.strip() and n not in INHERITED
                                and n not in ALLOWED_FONTS):
                            run.font.name = "Aptos"
                            log(idx, "font",
                                '%s -> Aptos on "%s"' % (n, run.text[:30]))

            # ---- whitespace -------------------------------------------
            # Scan runs in order and never merge them. Merging a
            # paragraph's runs destroys superscripts, inline bold and
            # symbol runs -- a real deck lost its TM that way.
            if want("whitespace") and sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    runs = para.runs
                    if not runs:
                        continue
                    full = "".join(r.text for r in runs)
                    if not full.strip():
                        continue
                    long_enough = len(full) > 25
                    last_idx = max(i for i, r in enumerate(runs) if r.text)
                    prev_space = False
                    touched = False
                    for i, run in enumerate(runs):
                        orig = run.text
                        if not orig:
                            continue
                        new_t = orig
                        if prev_space:
                            new_t = new_t.lstrip(" ")
                        if long_enough:
                            new_t = re.sub(r"  +", " ", new_t)
                        if i == last_idx:
                            new_t = new_t.rstrip()
                        if new_t != orig:
                            run.text = new_t
                            touched = True
                        if new_t:
                            prev_space = new_t.endswith(" ")
                    if touched:
                        log(idx, "whitespace",
                            'Spacing cleaned in "%s"' % full.strip()[:40])

            # ---- empty text box -----------------------------------------
            if (want("empty_box") and sh.shape_type == 17
                    and sh.has_text_frame
                    and not sh.text_frame.text.strip()
                    and not sh.is_placeholder):
                sh._element.getparent().remove(sh._element)
                log(idx, "empty_box",
                    "Removed empty text box at %.2f/%.2f" % (l or 0, t or 0))
                continue

            # ---- text overhang ------------------------------------------
            if (depth == 0 and want("text_overhang") and sh.has_text_frame
                    and sh.shape_type != 13 and None not in (l, t, w, h)):
                right = l + w
                if right > W + 0.02 and l >= 0:
                    new_w = max(W - EDGE_PAD - l, 1.0)
                    sh.width = Inches(new_w)
                    log(idx, "text_overhang",
                        "Width %.2f -> %.2f so the right edge sits inside "
                        "the slide" % (w, new_w))

            # ---- title drift ---------------------------------------------
            if want("title_drift") and dom and role == "content":
                try:
                    if (sh.is_placeholder
                            and "TITLE" in str(sh.placeholder_format.type)
                            and (abs(l - dom[0]) > POS_TOL
                                 or abs(t - dom[1]) > POS_TOL)):
                        sh.left, sh.top = Inches(dom[0]), Inches(dom[1])
                        log(idx, "title_drift",
                            "Title %.2f/%.2f -> %.2f/%.2f"
                            % (l, t, dom[0], dom[1]))
                except Exception:
                    pass

        # ---- add missing chrome -----------------------------------------
        if wants_chrome and not has_foot and foot_tpl is not None \
                and want("footer_missing"):
            new = clone_onto(slide, foot_tpl)
            for sh in slide.shapes:
                if sh._element is new:
                    sh.left, sh.top = Inches(FOOTER_LEFT), Inches(FOOTER_TOP)
                    sh.width, sh.height = foot_tpl.width, foot_tpl.height
                    if sh.text_frame.text.strip() != footer_text:
                        set_text_keep_format(sh, footer_text)
            log(idx, "footer_missing", "Added footer")

        if wants_chrome and not has_num and num_tpl is not None \
                and want("pagenum_missing"):
            new = clone_onto(slide, num_tpl)
            for sh in slide.shapes:
                if sh._element is new:
                    sh.left, sh.top = Inches(PAGENUM_LEFT), Inches(PAGENUM_TOP)
                    sh.width, sh.height = num_tpl.width, num_tpl.height
            log(idx, "pagenum_missing", "Added page-number field")

    prs.save(a.out)

    if a.log:
        with open(a.log, "w") as f:
            f.write("# Change log\n\n")
            f.write("Source: %s\nOutput: %s\nFooter: %s\n\n"
                    % (a.deck, a.out, footer_text))
            f.write("%d change(s).\n\n" % len(changes))
            cur = None
            for slide, kind, msg in sorted(
                    changes, key=lambda c: (c[0] or 0)):
                if slide != cur:
                    f.write("\n## Slide %s\n\n" % slide)
                    cur = slide
                f.write("- [%s] %s\n" % (kind, msg))

    print("%d change(s) -> %s" % (len(changes), a.out))
    for slide, kind, msg in changes[:200]:
        print("  s%-4s %-16s %s" % (slide, kind, msg[:80]))


if __name__ == "__main__":
    main()
